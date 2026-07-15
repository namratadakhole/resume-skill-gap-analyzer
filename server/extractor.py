import io
import pypdf
import docx

def extract_text_from_pdf(file_bytes: bytes) -> str:
    """
    Extracts text from PDF file bytes using pypdf.
    """
    try:
        pdf_file = io.BytesIO(file_bytes)
        reader = pypdf.PdfReader(pdf_file)
        text_content = []
        total_pages = len(reader.pages)
        print(f"[EXTRACT] Initialized PDF parser. Total pages: {total_pages}")
        
        for page_num in range(total_pages):
            try:
                page = reader.pages[page_num]
                page_text = page.extract_text()
                if page_text and page_text.strip():
                    text_content.append(page_text)
                    print(f"[EXTRACT] PDF Page {page_num + 1}/{total_pages}: Extracted {len(page_text)} chars")
                else:
                    print(f"[EXTRACT] PDF Page {page_num + 1}/{total_pages}: No text found (could be scanned/image)")
            except Exception as pe:
                print(f"[EXTRACT WARNING] Failed to parse PDF Page {page_num + 1}: {str(pe)}")
                continue
                
        extracted_text = "\n".join(text_content).strip()
        print(f"[EXTRACT] PDF parsing finished. Combined length: {len(extracted_text)} characters")
        if not extracted_text:
            raise ValueError(
                "No readable text found in the uploaded PDF document. "
                "If it is a scanned file or image, please upload the resume in PNG/JPG format or "
                "ensure it contains machine-readable text layers."
            )
        return extracted_text
    except ValueError as ve:
        raise ve
    except Exception as e:
        print(f"[EXTRACT ERROR] Failed to initialize PDF Reader: {str(e)}")
        import traceback
        traceback.print_exc()
        raise ValueError(f"Failed to extract text from PDF: {str(e)}")

def extract_text_from_docx(file_bytes: bytes) -> str:
    """
    Extracts text from DOCX file bytes using python-docx.
    """
    try:
        docx_file = io.BytesIO(file_bytes)
        doc = docx.Document(docx_file)
        text_content = []
        
        # Extract text from paragraphs
        paragraphs_count = len(doc.paragraphs)
        print(f"[EXTRACT] Initialized DOCX parser. Total paragraphs: {paragraphs_count}")
        for idx, para in enumerate(doc.paragraphs):
            try:
                if para.text and para.text.strip():
                    text_content.append(para.text)
            except Exception as pe:
                print(f"[EXTRACT WARNING] Failed to parse DOCX Paragraph {idx + 1}: {str(pe)}")
                continue
                
        # Extract text from tables
        tables_count = len(doc.tables)
        print(f"[EXTRACT] DOCX Tables found: {tables_count}")
        for idx, table in enumerate(doc.tables):
            try:
                for r_idx, row in enumerate(table.rows):
                    for c_idx, cell in enumerate(row.cells):
                        if cell.text and cell.text.strip():
                            text_content.append(cell.text)
            except Exception as te:
                print(f"[EXTRACT WARNING] Failed to parse DOCX Table {idx + 1}: {str(te)}")
                continue
                            
        extracted_text = "\n".join(text_content).strip()
        print(f"[EXTRACT] DOCX parsing finished. Combined length: {len(extracted_text)} characters")
        if not extracted_text:
            raise ValueError(
                "No readable text found in the uploaded DOCX document. "
                "Please verify the file contains standard paragraph text."
            )
        return extracted_text
    except ValueError as ve:
        raise ve
    except Exception as e:
        print(f"[EXTRACT ERROR] Failed to parse DOCX file: {str(e)}")
        import traceback
        traceback.print_exc()
        raise ValueError(f"Failed to extract text from DOCX: {str(e)}")

def extract_text_from_txt(file_bytes: bytes) -> str:
    """
    Extracts text from TXT file bytes. Tries multiple encodings.
    """
    encodings = ['utf-8', 'latin-1', 'cp1252', 'utf-16']
    for encoding in encodings:
        try:
            extracted_text = file_bytes.decode(encoding).strip()
            print(f"[EXTRACT] TXT decoding success using '{encoding}'. Length: {len(extracted_text)} characters")
            if not extracted_text:
                raise ValueError("No readable text found in the uploaded TXT document. The file is empty.")
            return extracted_text
        except UnicodeDecodeError:
            continue
    raise ValueError("Failed to decode TXT file. Please verify it is a valid text file with standard encoding.")

def extract_text_from_image(file_bytes: bytes, filename: str) -> str:
    """
    Extracts text from PNG, JPG, or JPEG file bytes using pytesseract and Pillow.
    """
    try:
        from PIL import Image
        import pytesseract
        import io
        
        # Load image using Pillow
        image = Image.open(io.BytesIO(file_bytes))
        
        # Run OCR
        extracted = pytesseract.image_to_string(image)
        
        cleaned = extracted.strip()
        if cleaned:
            return cleaned
        raise ValueError("No readable text found in the uploaded image.")
    except Exception as e:
        # Handle cases where tesseract CLI is missing or throws error
        err_msg = str(e)
        if "not installed" in err_msg or "PATH" in err_msg or "tesseract" in err_msg.lower():
            # Fallback to ocr.space free API as a robust hybrid fallback
            try:
                import requests
                payload = {
                    "apikey": "helloworld",
                    "language": "eng",
                    "isOverlayRequired": False
                }
                files = {
                    "file": (filename, file_bytes)
                }
                response = requests.post("https://api.ocr.space/parse/image", data=payload, files=files, timeout=20)
                if response.status_code == 200:
                    result = response.json()
                    exit_code = str(result.get("OCRExitCode", ""))
                    if exit_code == "1":
                        parsed_results = result.get("ParsedResults", [])
                        text_runs = [item.get("ParsedText", "") for item in parsed_results if item.get("ParsedText")]
                        extracted_web = "\n".join(text_runs).strip()
                        if extracted_web:
                            return extracted_web
            except Exception:
                pass
        raise ValueError("No readable text found in the uploaded image.")

def extract_text(file_bytes: bytes, filename: str) -> str:
    """
    Routes the file content to the appropriate extractor based on the filename extension.
    """
    lower_filename = filename.lower()
    if lower_filename.endswith('.pdf'):
        return extract_text_from_pdf(file_bytes)
    elif lower_filename.endswith('.docx'):
        return extract_text_from_docx(file_bytes)
    elif lower_filename.endswith(('.txt', '.md', '.rtf')):
        return extract_text_from_txt(file_bytes)
    elif lower_filename.endswith(('.png', '.jpg', '.jpeg', '.bmp', '.gif', '.tiff')):
        return extract_text_from_image(file_bytes, filename)
    else:
        raise ValueError(f"Unsupported file format: {filename}. Supported formats are PDF, DOCX, TXT, and Images (PNG, JPG, JPEG).")
